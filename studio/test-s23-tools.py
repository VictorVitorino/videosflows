"""Ferramentas do test-s23-pptx.js (não é bateria; o qa-gate só roda test-s[0-9][0-9]*.js).
  python3 test-s23-tools.py fonts <dir>      -> TTFs estáticos (Inter, Roboto, Roboto Condensed, JetBrains Mono: Regular/Bold, Roboto Light)
                                                 tirados dos woff2 variáveis de ../fonts2 + fonts.conf; imprime o caminho do fonts.conf
                                                 (simula “as fontes estão instaladas” para o LibreOffice). Precisa de fontTools + brotli.
  python3 test-s23-tools.py inspect <pptx>   -> JSON: python-pptx (slides, tamanho, ocultos, anotações, tipos de forma, textos, fontes),
                                                 zipfile.testzip, XML bem formado em todas as partes, tipos de conteúdo e relações.
"""
import sys, os, json, re, zipfile

def fonts(out):
    out = os.path.abspath(out)
    from fontTools.ttLib import TTFont
    from fontTools.varLib.instancer import instantiateVariableFont
    here = os.path.dirname(os.path.abspath(__file__)); src = os.environ.get('AM_FONTS_DIR') or os.path.join(here, '..', 'fonts2')
    css = open(os.path.join(src, 'gf.css'), encoding='utf-8').read()
    files = {}
    for m in re.finditer(r'/\* ([\w-]+) \*/\s*@font-face \{([^}]*)\}', css):
        if m.group(1) != 'latin': continue
        b = m.group(2); fam = re.search(r"font-family: '([^']+)'", b).group(1); url = re.search(r'url\(([^)]+)\)', b).group(1).split('/')[-1]
        files[fam] = url
    os.makedirs(out, exist_ok=True)
    want = [('Inter', 400, 'Regular', 'Inter'), ('Inter', 700, 'Bold', 'Inter'), ('Roboto', 400, 'Regular', 'Roboto'), ('Roboto', 700, 'Bold', 'Roboto'),
            ('Roboto', 300, 'Regular', 'Roboto Light'), ('Roboto Condensed', 400, 'Regular', 'Roboto Condensed'), ('Roboto Condensed', 700, 'Bold', 'Roboto Condensed'),
            ('JetBrains Mono', 400, 'Regular', 'JetBrains Mono'), ('JetBrains Mono', 700, 'Bold', 'JetBrains Mono')]
    made = []
    for fam, wt, sub, legacy in want:
        if fam not in files: continue
        dst = os.path.join(out, (legacy + '-' + sub).replace(' ', '') + '.ttf')
        if not os.path.exists(dst):
            f = TTFont(os.path.join(src, files[fam]))
            if 'fvar' in f:
                ax = {a.axisTag: a for a in f['fvar'].axes}
                loc = {t: (min(max(wt, a.minValue), a.maxValue) if t == 'wght' else a.defaultValue) for t, a in ax.items()}
                f = instantiateVariableFont(f, loc)
            f.flavor = None
            nm = f['name']
            for rec in list(nm.names):
                if rec.nameID in (1, 2, 4, 6, 16, 17): nm.removeNames(nameID=rec.nameID)
            full = legacy + ('' if sub == 'Regular' else ' ' + sub)
            nm.setName(legacy, 1, 3, 1, 0x409); nm.setName(sub, 2, 3, 1, 0x409); nm.setName(full, 4, 3, 1, 0x409); nm.setName(full.replace(' ', '') , 6, 3, 1, 0x409)
            if legacy != fam: nm.setName(fam, 16, 3, 1, 0x409); nm.setName('Light', 17, 3, 1, 0x409)
            os2 = f['OS/2']; os2.usWeightClass = wt
            os2.fsSelection = (os2.fsSelection & ~0b1100001) | (0b100000 if sub == 'Bold' else 0b1000000)
            f['head'].macStyle = 1 if sub == 'Bold' else 0
            f.save(dst)
        made.append(dst)
    conf = os.path.join(out, 'fonts.conf')
    open(conf, 'w').write('<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "fonts.dtd">\n<fontconfig><dir>%s</dir><include ignore_missing="yes">/etc/fonts/fonts.conf</include><cachedir>%s</cachedir></fontconfig>\n' % (out, os.path.join(out, 'cache')))
    print(json.dumps({'conf': conf, 'fonts': [os.path.basename(x) for x in made]}))

def inspect(path):
    import xml.etree.ElementTree as ET
    out = {'file': os.path.basename(path), 'size': os.path.getsize(path)}
    z = zipfile.ZipFile(path)
    out['zipBad'] = z.testzip()
    names = z.namelist(); out['parts'] = len(names); out['first'] = names[0]
    bad = []
    for n in names:
        if n.endswith('.xml') or n.endswith('.rels'):
            try: ET.fromstring(z.read(n))
            except Exception as e: bad.append(n + ': ' + str(e))
    out['xmlBad'] = bad
    ct = z.read('[Content_Types].xml').decode('utf-8')
    out['ctMissing'] = [n for n in names if not n.endswith('/') and n != '[Content_Types].xml' and ('/' + n) not in ct and n.rsplit('.', 1)[-1] not in re.findall(r'Extension="(\w+)"', ct)]
    relBad = []
    for n in names:
        if not n.endswith('.rels'): continue
        base = n.replace('_rels/', '').rsplit('.rels', 1)[0]; d = os.path.dirname(base)
        for t in re.findall(r'Target="([^"]+)"', z.read(n).decode('utf-8')):
            p = os.path.normpath(os.path.join(d, t)).replace('\\', '/')
            if p not in names: relBad.append(n + ' -> ' + t)
    out['relBad'] = relBad
    out['mediaKB'] = round(sum(z.getinfo(n).file_size for n in names if n.startswith('ppt/media/')) / 1024)
    from pptx import Presentation
    from pptx.util import Emu
    pr = Presentation(path)
    out['title'] = pr.core_properties.title
    out['w'] = pr.slide_width; out['h'] = pr.slide_height
    sl = []
    for s in pr.slides:
        el = s._element
        info = {'hidden': el.get('show') == '0', 'notes': s.notes_slide.notes_text_frame.text if s.has_notes_slide else None, 'shapes': []}
        def walk(shapes, depth):
            for sh in shapes:
                k = sh.shape_type
                d = {'kind': str(k).split('.')[-1].split(' ')[0] if k is not None else sh._element.tag.split('}')[-1], 'name': sh.name,
                     'x': round(Emu(sh.left).pt / .75, 1) if sh.left is not None else None, 'y': round(Emu(sh.top).pt / .75, 1) if sh.top is not None else None,
                     'w': round(Emu(sh.width).pt / .75, 1) if sh.width is not None else None, 'h': round(Emu(sh.height).pt / .75, 1) if sh.height is not None else None,
                     'rot': round(sh.rotation, 2) if hasattr(sh, 'rotation') else 0, 'tag': sh._element.tag.split('}')[-1]}
                x = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm')
                if x is not None: d['flipH'] = x.get('flipH') == '1'; d['flipV'] = x.get('flipV') == '1'
                g = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
                if g is not None: d['prst'] = g.get('prst')
                if getattr(sh, 'has_text_frame', False) and sh.has_text_frame and sh.text_frame.text.strip():
                    d['text'] = sh.text_frame.text
                    runs = []
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            f = r.font
                            runs.append({'t': r.text, 'b': f.bold, 'i': f.italic, 'u': bool(f.underline), 'sz': f.size.pt if f.size else None, 'font': f.name,
                                         'color': str(f.color.rgb) if f.color and f.color.type is not None else None})
                    d['runs'] = runs; d['paras'] = len(sh.text_frame.paragraphs)
                if d['tag'] == 'cxnSp':
                    ln = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
                    he = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}headEnd'); te = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd')
                    d['head'] = he.get('type') if he is not None else None; d['tail'] = te.get('type') if te is not None else None
                    d['dash'] = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}custDash') is not None
                if d['tag'] == 'pic':
                    sr = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect')
                    d['crop'] = dict(sr.attrib) if sr is not None else None
                    am = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}alphaModFix')
                    d['alpha'] = int(am.get('amt')) / 100000 if am is not None else None
                    try: d['img'] = sh.image.content_type
                    except Exception: d['img'] = None
                info['shapes'].append(d)
                if d['tag'] == 'grpSp': walk(sh.shapes, depth + 1)
        walk(s.shapes, 0)
        bg = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
        c = bg.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr') if bg is not None else None
        info['bg'] = c.get('val') if c is not None else None
        sl.append(info)
    out['slides'] = sl
    print(json.dumps(out, ensure_ascii=False))

if __name__ == '__main__':
    {'fonts': fonts, 'inspect': inspect}[sys.argv[1]](sys.argv[2])
