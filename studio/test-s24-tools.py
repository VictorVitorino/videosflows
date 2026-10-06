"""Ferramentas do test-s24-import.js (não é bateria; o qa-gate só roda test-s[0-9][0-9]*.js).
  python3 test-s24-tools.py make <dir>   -> gera os arquivos de cobertura da importação com python-pptx:
      fx-a.pptx  16:9 — textos com formato por trecho, formas (retângulo, arredondado, elipse, chevron, seta, losango, triângulo),
                 conectores (reta com ponta, cotovelo), foto (com recorte), grupo, tabela, gráfico de colunas, slide oculto,
                 anotações do orador, marcadores, fonte/cor/tamanho mistos, giro e espelho
      fx-b.pptx  4:3 — espaços reservados (título + corpo com marcadores de vários níveis) herdando do layout/mestre, subtítulo
      fx-c.pptx  16:9 — um slide só com um texto grande e uma foto (teste do PDF)
  python3 test-s24-tools.py pdf <pptx> <outdir> -> converte com o LibreOffice (soffice) para PDF e imprime o caminho
"""
import sys, os, json, subprocess

def make(out):
    from pptx import Presentation
    from pptx.util import Emu, Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from PIL import Image, ImageDraw
    os.makedirs(out, exist_ok=True)
    png = os.path.join(out, 'foto.png'); im = Image.new('RGB', (400, 300), '#1B7F3B'); dr = ImageDraw.Draw(im); dr.rectangle([0, 0, 200, 300], fill='#F78C16'); dr.ellipse([120, 60, 280, 240], fill='#002A46'); im.save(png)

    # ---------- fx-a: 16:9 cobertura ----------
    prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    tb = s1.shapes.add_textbox(Inches(.5), Inches(.4), Inches(9), Inches(1)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = 'Receita '; r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = RGBColor(0x00, 0x2A, 0x46); r.font.name = 'Arial'
    r2 = p.add_run(); r2.text = 'cresce 18%'; r2.font.size = Pt(32); r2.font.bold = True; r2.font.italic = True; r2.font.color.rgb = RGBColor(0xF7, 0x8C, 0x16); r2.font.name = 'Arial'
    r3 = p.add_run(); r3.text = ' com margem — ação, coração'; r3.font.size = Pt(24); r3.font.color.rgb = RGBColor(0x3E, 0x4C, 0x5E); r3.font.name = 'Arial'
    p2 = tf.add_paragraph(); p2.text = 'Segunda linha, alinhada à direita e sublinhada'; p2.alignment = PP_ALIGN.RIGHT; p2.runs[0].font.underline = True; p2.runs[0].font.size = Pt(14); p2.runs[0].font.name = 'Calibri'
    # formas
    def shp(kind, x, y, w, h, fill, text=None, line=None, lw=None):
        sh = s1.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line is None: sh.line.fill.background()
        else: sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw or 1.5)
        if text:
            sh.text_frame.text = text; sh.text_frame.paragraphs[0].runs[0].font.size = Pt(14); sh.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); sh.text_frame.paragraphs[0].runs[0].font.bold = True
        return sh
    shp(MSO_SHAPE.RECTANGLE, .5, 1.8, 2.2, 1.3, '002A46', 'Retângulo')
    rr = shp(MSO_SHAPE.ROUNDED_RECTANGLE, 3.0, 1.8, 2.2, 1.3, '43698F', 'Arredondado', line='F78C16', lw=3)
    el = shp(MSO_SHAPE.OVAL, 5.5, 1.8, 1.4, 1.3, 'F78C16'); el.rotation = 30
    ch = shp(MSO_SHAPE.CHEVRON, 7.2, 1.8, 2.4, 1.1, '1B7F3B', 'Chevron')
    ar = shp(MSO_SHAPE.RIGHT_ARROW, 9.9, 1.8, 2.6, 1.1, 'C0392B', 'Seta'); ar._element.spPr.xfrm.set('flipH', '1')
    shp(MSO_SHAPE.DIAMOND, .5, 3.5, 1.3, 1.3, '7B5BB3'); shp(MSO_SHAPE.ISOSCELES_TRIANGLE, 2.0, 3.5, 1.3, 1.3, None, line='002A46', lw=2)
    shp(MSO_SHAPE.HEXAGON, 3.6, 3.5, 1.5, 1.3, 'EEF2F7', 'Hex', line='A3B8D6', lw=1)
    hx = s1.shapes[-1]; hx.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0x2A, 0x46)
    # conectores
    c1 = s1.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.5), Inches(3.8), Inches(8.5), Inches(4.6)); c1.line.color.rgb = RGBColor(0, 0x2A, 0x46); c1.line.width = Pt(3)
    ln = c1.line._get_or_add_ln(); from lxml import etree
    te = etree.SubElement(ln, '{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd'); te.set('type', 'triangle')
    c2 = s1.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(9), Inches(3.6), Inches(12), Inches(4.8)); c2.line.color.rgb = RGBColor(0xF7, 0x8C, 0x16); c2.line.width = Pt(2.25); c2.line.dash_style = 4  # dash
    # foto com recorte
    pic = s1.shapes.add_picture(png, Inches(.5), Inches(5.1), Inches(2.4), Inches(1.8)); pic.crop_left = .25; pic.crop_right = .0
    # grupo
    grp = s1.shapes.add_group_shape(); g1 = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(5.2), Inches(1.4), Inches(1.0)); g1.fill.solid(); g1.fill.fore_color.rgb = RGBColor(0xA3, 0xB8, 0xD6); g1.line.fill.background()
    g2 = grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.1), Inches(5.2), Inches(1.0), Inches(1.0)); g2.fill.solid(); g2.fill.fore_color.rgb = RGBColor(0x00, 0x2A, 0x46); g2.line.fill.background()
    # tabela
    tbl = s1.shapes.add_table(3, 3, Inches(6.6), Inches(5.1), Inches(5.9), Inches(1.6)).table
    for ci, t in enumerate(['Iniciativa', 'Prazo', 'Status']): tbl.cell(0, ci).text = t
    for ri, row in enumerate([['Compras', 'Q3', 'Em curso'], ['Logística', 'Q4', 'Planejado']], 1):
        for ci, t in enumerate(row): tbl.cell(ri, ci).text = t
    tbl.cell(1, 2).fill.solid(); tbl.cell(1, 2).fill.fore_color.rgb = RGBColor(0x1B, 0x7F, 0x3B)
    # notas
    s1.notes_slide.notes_text_frame.text = 'Abertura: falar da receita.\nSegunda linha das anotações.'
    # slide 2: gráfico + marcadores + texto girado
    s2 = prs.slides.add_slide(blank)
    cd = CategoryChartData(); cd.categories = ['2023', '2024', '2025']; cd.add_series('Varejo', (42, 48, 55)); cd.add_series('Digital', (8, 14, 22))
    gf = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.5), Inches(.5), Inches(6), Inches(4), cd); gf.chart.has_title = True; gf.chart.chart_title.text_frame.text = 'Receita por canal'
    tb2 = s2.shapes.add_textbox(Inches(7), Inches(.5), Inches(5.5), Inches(4)); tf2 = tb2.text_frame; tf2.word_wrap = True
    tf2.text = 'Primeiro ponto'
    for i, t in enumerate(['Segundo ponto', 'Subitem do segundo', 'Terceiro ponto']):
        pp = tf2.add_paragraph(); pp.text = t; pp.level = 1 if i == 1 else 0
    for pp in tf2.paragraphs:
        pPr = pp._p.get_or_add_pPr(); bu = etree.SubElement(pPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}buChar'); bu.set('char', '•'); pPr.set('marL', str(342900 * (pp.level + 1))); pPr.set('indent', '-342900')
        pp.runs[0].font.size = Pt(18); pp.runs[0].font.name = 'Calibri'
    rt = s2.shapes.add_textbox(Inches(.5), Inches(5), Inches(4), Inches(.8)); rt.text_frame.text = 'Texto girado 15°'; rt.rotation = 15; rt.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    pil = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7), Inches(5), Inches(5), Inches(1)); pil.adjustments[0] = .5; pil.fill.solid(); pil.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF7); pil.line.fill.background(); pil.text_frame.text = 'Pílula (roundRect 50%)'; pil.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0x2A, 0x46); pil.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    # slide 3 oculto
    s3 = prs.slides.add_slide(blank); s3._element.set('show', '0'); t3 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)); t3.text_frame.text = 'Slide oculto'
    prs.core_properties.title = 'Cobertura importação A'
    prs.save(os.path.join(out, 'fx-a.pptx'))

    # ---------- fx-b: 4:3 com espaços reservados ----------
    prs = Presentation(); prs.slide_width = Emu(9144000); prs.slide_height = Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = 'Título do projeto'; s.placeholders[1].text = 'Subtítulo herdado do layout'
    s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = 'Agenda'; body = s.placeholders[1].text_frame; body.text = 'Contexto e objetivos'
    for t, lv in [('Diagnóstico', 0), ('Entrevistas', 1), ('Dados', 1), ('Plano de ação', 0)]:
        pp = body.add_paragraph(); pp.text = t; pp.level = lv
    s.notes_slide.notes_text_frame.text = 'Falar da agenda.'
    prs.core_properties.title = 'Cobertura importação B'
    prs.save(os.path.join(out, 'fx-b.pptx'))

    # ---------- fx-c: texto + foto (para o PDF) ----------
    prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2)); bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0, 0x2A, 0x46); bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(.5), Inches(.25), Inches(10), Inches(.8)); t.text_frame.text = 'Programa de Formação'; t.text_frame.paragraphs[0].runs[0].font.size = Pt(36); t.text_frame.paragraphs[0].runs[0].font.bold = True; t.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); t.text_frame.paragraphs[0].runs[0].font.name = 'Arial'
    b = s.shapes.add_textbox(Inches(.5), Inches(1.6), Inches(7), Inches(3)); b.text_frame.word_wrap = True
    b.text_frame.text = 'Objetivo: preparar a equipe para a nova estrutura.'
    pp = b.text_frame.add_paragraph(); pp.text = 'Público: 120 pessoas em três ondas.'
    pp = b.text_frame.add_paragraph(); pp.text = 'Prazo: setembro a dezembro.'
    for pp in b.text_frame.paragraphs: pp.runs[0].font.size = Pt(20); pp.runs[0].font.color.rgb = RGBColor(0x3E, 0x4C, 0x5E); pp.runs[0].font.name = 'Arial'
    c = s.shapes.add_textbox(Inches(.5), Inches(5.8), Inches(12), Inches(.6)); c.text_frame.text = 'Centralizado em laranja'; c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER; c.text_frame.paragraphs[0].runs[0].font.size = Pt(18); c.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xF7, 0x8C, 0x16); c.text_frame.paragraphs[0].runs[0].font.name = 'Arial'
    s.shapes.add_picture(png, Inches(8.5), Inches(1.8), Inches(4), Inches(3))
    prs.save(os.path.join(out, 'fx-c.pptx'))

    # ---------- fx-big: 40 slides (cancelar no meio) ----------
    prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    for i in range(40):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        t = s.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(1)); t.text_frame.text = 'Slide %d de 40' % (i + 1)
        for k in range(6):
            sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.5 + k * 2), Inches(2), Inches(1.8), Inches(1.2)); sh.text_frame.text = 'Card %d' % (k + 1)
        s.shapes.add_picture(png, Inches(.5), Inches(4), Inches(3), Inches(2.2))
    prs.save(os.path.join(out, 'fx-big.pptx'))
    print(json.dumps({'dir': out, 'files': ['fx-a.pptx', 'fx-b.pptx', 'fx-c.pptx', 'fx-big.pptx']}))

def pdf(src, out):
    os.makedirs(out, exist_ok=True)
    env = dict(os.environ); env['HOME'] = out
    r = subprocess.run(['soffice', '--headless', '-env:UserInstallation=file://' + os.path.abspath(out) + '/.lo', '--convert-to', 'pdf', '--outdir', out, src], capture_output=True, text=True, timeout=240, env=env)
    dst = os.path.join(out, os.path.splitext(os.path.basename(src))[0] + '.pdf')
    print(json.dumps({'pdf': dst if os.path.exists(dst) else None, 'err': r.stderr[-300:]}))

if __name__ == '__main__':
    {'make': lambda: make(sys.argv[2]), 'pdf': lambda: pdf(sys.argv[2], sys.argv[3])}[sys.argv[1]]()
